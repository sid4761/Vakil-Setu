import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('C:/Users/DELL/Vakil-Setu')
BUILD = ROOT / '.codex-deck/project-framework'
OUT = ROOT / 'output/project-presentation'
slides = json.loads((BUILD / 'content.json').read_text(encoding='utf-8'))['slides']
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.78)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Arial'
styles['Normal'].font.size = Pt(11.5)
styles['Normal'].paragraph_format.line_spacing = 1.12
styles['Normal'].paragraph_format.space_after = Pt(8)
styles['Title'].font.size = Pt(29)
styles['Title'].font.color.rgb = RGBColor.from_string('000000')
styles['Heading 1'].font.size = Pt(19)
styles['Heading 1'].font.color.rgb = RGBColor.from_string('B12632')
styles['Heading 1'].paragraph_format.space_before = Pt(12)
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].font.color.rgb = RGBColor.from_string('172A35')
styles['Heading 2'].paragraph_format.space_before = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)

def para(text, style=None):
    return doc.add_paragraph(text, style)

def note(label, text):
    p=doc.add_paragraph()
    r=p.add_run(label+' ');r.bold=True
    p.add_run(text)
    p.paragraph_format.space_after=Pt(9)

doc.add_heading('Vakil Setu Presentation Script',0)
para('Faculty presentation and viva preparation', 'Subtitle')
para('About 10 minutes for 14 slides. Add about 90 seconds if demonstrating the application. The speaking text also appears in the PowerPoint speaker notes.')
doc.add_heading('Presentation scope',1)
para('This presentation explains the maintained Next.js, Convex and Clerk rebuild. It distinguishes implemented workflows from recorded runtime evidence and remaining acceptance work. The legacy React and Express application remains separate.')
doc.add_heading('Before presenting',1)
for t in [
    'Confirm each member\'s real contribution and assign speaking sections. The contribution slide names workstreams without assigning personal ownership.',
    'Rehearse with the installed presentation software. Use the notes as spoken paragraphs rather than reading the slide text.',
    'Prepare a labelled synthetic example. Use no real case, identity document or private client details in the demonstration.',
    'Check the development application before presenting. If a service is unavailable, explain the state and use the screenshot and recorded evidence.',
    'Keep these boundaries explicit: no verified DGX inference, no measured legal accuracy, no claim of complete support for 23 languages, and no guarantee of emergency connection.'
]: para(t,'List Bullet')
doc.add_heading('Suggested speaking allocation',1)
para('For one presenter, deliver all sections. For a group, assign slides 1–4, 5–9 and 10–14 to available members. This is a speaking allocation, not a claim about who developed the components.')
note('Timing', 'The slide targets total about nine minutes of speech, leaving roughly one minute for transitions. Adjust to your natural speaking pace.')

for i,d in enumerate(slides):
    if i % 2 == 0: doc.add_page_break()
    doc.add_heading(f"Slide {i+1:02d} {d['title']}",1)
    note('Section',d['section'].replace('  /  ',' — ')+'    |    Target '+str(d['seconds'])+' seconds')
    para(d['script'])
    note('Presenter cue',d['cue'])
    if i%2==0: para('')

doc.add_page_break()
doc.add_heading('Optional live demonstration',0)
para('Use this after slide 7 or after the main presentation. Target about 90 seconds. Complete the save operation only in an authorised development account.')
steps=[
('Begin on the landing page', '“I will use a fictional example to show the intake and review flow.”'),
('Open intake and select a pilot language', '“The interface language and the intake language are recorded separately. Typed input remains available if speech is unavailable.”'),
('Enter the labelled example below', '“These are the user\'s words. We can correct them before saving.”'),
('Acknowledge review and save privately', '“Saving records the intake. It does not automatically generate AI analysis.”'),
('Open the manual summary editor', '“The starter uses the saved intake. The user must review the summary before saving.”'),
('Explain the next step without submitting a booking', '“The user can browse approved advocates. A consultation request needs a future slot, specific sharing consent and advocate acceptance.”')
]
for title, words in steps:
    doc.add_heading(title,2);para(words)
doc.add_heading('Synthetic example',2)
para('[TEST] Deposit intake demo\nI moved out of a rented room on 1 October. My landlord has not returned my deposit of Rs 10,000. I have the agreement and payment receipt. I want to speak with an advocate about my options.')
note('Fallback', 'If login or persistence fails, stop the submission, show the slide screenshot, and describe the recorded synthetic Punjabi save-and-reload result. Never announce that a failed save succeeded.')

doc.add_page_break()
doc.add_heading('Individual contribution preparation',0)
para('Each member should prepare one truthful paragraph using the pattern below. Replace the blanks with work that the member actually performed and can explain during the viva.')
para('“My specific responsibility was _____. I personally handled _____ in the _____ component. I contributed to the decision to _____ because _____. The main challenge I addressed was _____. I used _____ to resolve it, and I checked the result by _____. The outcome of my work was _____. The remaining limitation is _____.”')
doc.add_heading('Evidence to prepare',1)
for t in [
    'One component or backend operation you can walk through.',
    'One design choice and the alternative you considered.',
    'One concrete failure or edge case you addressed.',
    'One check demonstrating the resulting behaviour.',
    'A clear account of code, tools or assistance you used.'
]: para(t,'List Bullet')
doc.add_heading('Questions the PROJECT structure answers',1)
mapping=[
('Problem', 'Slides 2–3 identify the people affected, friction, existing approaches and the project gap.'),
('Reason', 'Slide 3 explains relevance, intended beneficiaries and established alternatives.'),
('Objective', 'Slide 4 defines features, boundaries and completion criteria.'),
('Journey', 'Slide 5 explains the audit, requirements, alternative scope and verification approach.'),
('Execution', 'Slides 6–9 explain tools, component boundaries, intake, consultation and capability status.'),
('Contribution', 'Slide 10 and the paragraph above cover role, task, decisions, challenge, skills and outcome.'),
('Takeaway', 'Slides 11–14 explain results, partial objective achievement, evaluation, learning, limits and improvements.')
]
for title,body in mapping: note(title,body)

viva=[
('What problem does Vakil Setu solve?', 'It addresses the first interaction: describing a legal concern, reviewing a private summary and requesting a consultation with a chosen advocate. Whether it improves user outcomes remains a pilot hypothesis.'),
('Why build this when Tele Law and LawRato exist?', 'Those services show that digital access to lawyers is established. Our project studies the intake and review process before a consultation, including specific-version sharing consent. We have not established comparative superiority.'),
('What is the current stack?', 'Next.js and React provide the web interface. TypeScript supports shared types. Clerk handles sessions and Convex handles authorised application operations and data. Bun and Turborepo organise the workspace. The legacy React, Express and MongoDB stack is separate.'),
('Is the AI running on DGX Spark?', 'That is unverified. The gateway adapter is implemented, but the endpoint contract and verified inference evidence are absent. A gateway name or model label does not prove hardware attribution.'),
('Does saving a case run AI?', 'No. Saving stores the reviewed intake privately. The user can start a manual summary. A provider draft requires configuration, consent and review.'),
('How do you prevent AI hallucinations?', 'We do not claim to eliminate them. The adapter validates structure, locale and output bounds, which can reject malformed results. Those checks cannot establish legal correctness. Legal review and a factual evaluation set are still required.'),
('How many languages are supported?', 'English, Hindi and Punjabi are the pilot interfaces, with some English fallbacks and incomplete human review. The registry contains 22 scheduled languages plus English as a broader target. Speech and AI capabilities have separate evidence boundaries.'),
('What stops another user reading a private case?', 'Backend operations derive the authenticated account and check record ownership or an explicitly permitted relationship. Hiding UI controls is insufficient. In-memory tests exercise access rules, while independent-account hosted checks still need completion.'),
('How do you handle two people requesting one slot?', 'The request mutation checks and reserves the slot atomically and rejects unavailable slots. An idempotency key also prevents duplicate request records for repeated identical submissions. Competing reservations have automated test coverage. Hosted concurrency acceptance remains pending.'),
('What exactly do the 78 tests prove?', 'The recorded total is 60 backend, 12 language and six case-render tests. They exercise specified behaviours with synthetic boundaries where documented. They do not prove actual speech accuracy, human comprehension, live AI quality or every hosted workflow.'),
('Is the application production ready?', 'No. The record supports local workflows and a limited authenticated hosted intake check. Cross-account and full booking acceptance, language review, operational policies and speech-asset provenance review remain release work.'),
('What does it cost per consultation?', 'There is no measured unit-cost model yet. A pilot should record hosting, storage, speech and AI usage plus support and credential-review time, then divide total relevant costs by completed consultations. A successful build says nothing about affordable unit economics.'),
('What would you improve first?', 'I would finish hosted multi-role acceptance and run a small observed pilot before adding features. I would measure task completion, corrections to names, dates and amounts, sharing-consent comprehension and booking conflicts.'),
('What did you personally contribute?', 'Answer with the confirmed contribution paragraph. Name one component, one decision, one challenge and the check you performed. Do not claim sole authorship of inherited code or work performed by another team member.')
]
for start in [0,5,10]:
    doc.add_page_break()
    doc.add_heading('Viva questions and answers',0)
    for idx,(q,a) in enumerate(viva[start:start+5],start+1):
        doc.add_heading(f'{idx:02d} '+q.rstrip('?'),2);para(a)

doc.add_page_break()
doc.add_heading('Source and claim notes',0)
para('Project status evidence dates are 14 to 16 October 2026. The deck reports recorded results and does not represent a fresh execution of the project tests.')
doc.add_heading('Primary project records',1)
for s in ['rebuild/README.md','rebuild/docs/requirements.md','rebuild/docs/decisions.md','rebuild/docs/legacy-ai-audit.md','rebuild/docs/verification.md','rebuild/docs/hosted-activation.md','rebuild/docs/language-matrix.md','rebuild/packages/backend/convex/bookings.ts','rebuild/packages/backend/convex/lib/provider.ts']:
    para(s)
doc.add_heading('External service context',1)
para('Tele-Law service description\nhttps://www.tele-law.in/faq.html')
para('LawRato service description\nhttps://lawrato.com/cmses/page/how-does-lawrato-work')
para('Accessed 17 October 2026. These sources support the service descriptions on slide 3. The presentation does not claim either service lacks the capabilities explored by Vakil Setu.')
doc.add_heading('Claims to avoid',1)
para('Do not claim zero hallucinations, zero downtime, verified DGX hardware, legally certified outputs, an invented accuracy percentage, full support for 23 languages, a completed consultation from a pending request, or personal ownership that the team has not confirmed.')
OUT.mkdir(parents=True,exist_ok=True)
target=OUT/'Vakil_Setu_Presentation_Script.docx'
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
styles['Subtitle'].font.color.rgb = RGBColor.from_string('52616A')
styles['Subtitle'].font.italic = False
doc.save(target)
print(target)
print('Speaking words:',sum(len(s['script'].split()) for s in slides))

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\DELL\Vakil-Setu')
OUT=ROOT/'output/reports/Vakil_Setu_SRS_From_Scratch.docx'
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing=1.1
d.styles['Title'].font.size=Pt(28)
d.styles['Heading 1'].font.size=Pt(20)
d.styles['Heading 2'].font.size=Pt(12.5)
for name in ['Heading 1','Heading 2','Heading 3']:
    d.styles[name].paragraph_format.keep_with_next=True
    d.styles[name].paragraph_format.space_before=Pt(8)
hp=sec.header.paragraphs[0];hp.text='VAKIL SETU    SOFTWARE REQUIREMENTS SPECIFICATION'
hp.style='Normal';hp.runs[0].font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('VS SRS 001   /   Draft 1.0   /   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
for r in fp.runs:r.font.size=Pt(8)
d.core_properties.title='Vakil Setu Software Requirements Specification'
d.core_properties.subject='From scratch working MVP specification'
d.core_properties.author='Raghav Nargotra and Sidarth Gupta'

def p(t,style=None):
    q=d.add_paragraph(t,style);q.paragraph_format.widow_control=True;return q
def h(t,l=2):d.add_heading(t,l)
def page(t):d.add_page_break();h(t,1)
def req(code,t):
    q=d.add_paragraph();q.add_run(code+'  ').bold=True;q.add_run(t);q.paragraph_format.space_after=Pt(6)
def table(head,rows,widths=None,size=9.5):
    t=d.add_table(rows=1, cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,v in enumerate(head):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trPr.append(ns)
        if ri==0:
            hdr=OxmlElement('w:tblHeader');trPr.append(hdr)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17365D' if ri==0 else ('F1F5FA' if ri%2 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(1);q.paragraph_format.line_spacing=1.05
                for r in q.runs:r.font.size=Pt(size);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000');r.bold=ri==0
    p('');return t

# 1
p('Vakil Setu', 'Title')
p('Software Requirements Specification','Subtitle')
p('From scratch development of an Indian multilingual legal intake and advocate consultation platform')
table(['Document control','Value'],[
 ['Identifier','VS SRS 001'],['Version and status','1.0  •  Draft for development review'],['Date','14 October 2026'],
 ['Project team','Raghav Nargotra 2023A1R033\nSidarth Gupta 2023A1R038'],['Development approach','New implementation with no dependency on the previous repository']], [1.65,5.18],10)
h('Purpose and delivery objective')
p('This specification defines the first working release of Vakil Setu for developers, project evaluators, and pilot reviewers. It is a requirements baseline for a new build. It does not describe completed functionality, certify legal compliance, or depend on previous source code.')
p('The release succeeds when two independent accounts complete a persistent journey: user registration, multilingual issue intake, suitable advocate discovery, an available slot request, advocate acceptance, and visible booking status. An administrator must approve the advocate before public discovery.')
h('Document navigation')
p('1 Purpose and scope   •   2 Users and workflows   •   3 Architecture\n4 Accounts and advocate onboarding   •   5 Multilingual intake and triage\n6 Discovery and private summaries   •   7 Slots and bookings\n8 Interface contracts   •   9 Data requirements   •   10 Quality and security\n11 Acceptance and traceability   •   12 Delivery and open decisions')
p('Requirement language: SHALL identifies mandatory behavior. SHOULD identifies a recommended improvement. MAY identifies an optional extension. All numbered requirements are proposed until this draft is approved.')

# 2
page('1 Purpose and scope')
h('Product purpose')
p('Vakil Setu helps a person explain a legal issue in a familiar language, understand preliminary information, and request a consultation with a suitable advocate. The user remains responsible for reviewing the transcript and choosing the professional. AI output is preliminary information and must be distinguishable from an advocate response.')
h('Release boundary')
table(['Included in the working MVP','Deferred beyond this release'],[
 ['User and advocate accounts; protected administrator approval','Payment collection, commissions, refunds and reconciliation'],
 ['Multilingual text, supported speech input and spoken playback','Public Q&A, reviews, rankings based on reviews and guide library'],
 ['Structured triage, follow-up context and explicit failure handling','Document OCR, statutory deadline calculation and legal drafting'],
 ['Approved advocate profiles, expertise/language/location filters','Native mobile applications, built-in video calls and court integrations'],
 ['Advocate slots, consultation requests and acceptance','Guaranteed emergency response and automatic advocate connection'],
 ['Private user-approved summary and booking dashboards','Automated mass scraping or third-party profile imports']], [3.42,3.41],10)
h('Indian multilingual scope')
p('The product scope is the 22 scheduled Indian languages, with English also supported. Language capability shall be tracked separately for interface text, typed input, AI output, speech recognition, and playback. The release must accurately identify which capabilities work for each language. Coverage is not established merely by displaying a language in a menu.')
p('For the time-constrained pilot, the proposed acceptance target is tested end-to-end interaction in at least two Indian languages plus English. The team must record the chosen pilot languages before implementation. Additional languages may remain text-capable or unavailable for a particular speech provider, with a clear fallback.')
h('Dependencies and assumptions')
p('The MVP requires a browser, network access, an application server, and a persistent database. Speech and AI services depend on configured provider capabilities. An authorized administrator reviews advocate applications. Sample data must remain visibly labelled and separate from real onboarding.')
h('Priority convention')
p('P0 requirements form the working journey and its security boundaries. P1 items may improve a later iteration but must not delay P0 integration. This SRS assigns every numbered requirement below to P0 unless explicitly marked otherwise.')

# 3
page('2 Users and workflows')
table(['Role','Permitted responsibilities','Access boundary'],[
 ['Visitor','Read public information and approved profiles; register','Cannot access private cases or create bookings'],
 ['User','Manage own account, intake, summary and bookings','Own private information and authorized advocate sharing only'],
 ['Advocate','Submit profile, maintain slots, accept or decline requests','Own profile and requests assigned to that advocate'],
 ['Administrator','Review applications and manage approval/suspension','Credential review; no unrestricted routine access to case content']], [1.1,3.13,2.6],10)
h('UC 01 User intake and discovery')
p('A registered user chooses a language, speaks or types an issue, reviews the text, and submits it. The system returns structured preliminary information and suitable approved profiles. The user can correct the suggested category, change filters, or ask a contextual follow-up. A failed speech attempt preserves typed input as an alternative.')
h('UC 02 Advocate onboarding')
p('An advocate registers, submits professional details and credentials, and sees a pending status. An administrator reviews the application and approves or rejects it with a reason. The profile appears publicly only after approval. A verification label shall state the actual review performed; it shall not imply guaranteed competence or outcomes.')
h('UC 03 Request and accept consultation')
p('An approved advocate publishes a future slot. A user selects it, approves the summary to share, and submits a consultation request. The system reserves the slot atomically. The advocate accepts or declines; the user sees the persisted decision. No appointment is called accepted until the advocate accepts it.')
h('UC 04 Conflict or cancellation')
p('When two users request the same slot, one succeeds and the other receives a conflict with refreshed availability. A user can cancel their own future request. An advocate can cancel an assigned future appointment with a reason. The system updates booking history and releases an eligible slot consistently.')
h('UC 05 Suspension or service failure')
p('Suspending an advocate hides the profile and prevents new requests. Existing affected bookings are flagged for administrator handling and their users are notified. If the AI provider fails, the application identifies the limited fallback or allows retry; it never presents a fabricated successful provider result.')

# 4
page('3 Architecture and system boundaries')
p('Build one responsive portal with role-specific views. A single backend owns authentication, permissions, matching, slot reservation, and booking state. External providers must not determine application ownership or write directly to booking records.')
table(['Layer','Recommended technology','Responsibilities'],[
 ['Presentation','React with Vite','Registration, role dashboards, multilingual intake, transcript review, profiles and booking status'],
 ['Application API','Node.js with Express','Authentication, validation, permissions, triage coordination, matching and state transitions'],
 ['Persistence','MongoDB with Mongoose','Accounts, applications, profiles, slots, private cases, bookings and audit history'],
 ['AI adapter','Configured provider behind a server adapter','Language-aware structured classification, timeout handling and limited fallback'],
 ['Speech adapter','Browser APIs or a configured provider','Capability discovery, recognition/playback and typed fallback'],
 ['Private file storage','Protected local development storage or configured object storage','Credential documents with controlled access and limits']], [1.12,2.0,3.71],10)
h('Request flow')
p('The browser sends reviewed text and a selected language to the authenticated triage API. The backend validates the request, retrieves permitted context, calls the AI adapter, validates the result, and queries approved advocate profiles. It returns the preliminary result, limitations, and matches. A booking request refers to a slot and a user-approved summary, rather than trusting client-supplied ownership or time information.')
h('Development boundaries')
p('Person 1 owns backend logic, persistence, migrations, and the API contract. Person 2 owns frontend flows, accessibility, client integration, and frontend tests. Agree on contracts before incompatible edits. Each developer may use configured coding agents, but file ownership and final integration responsibility remain explicit.')
h('External service boundaries')
p('Server-side secrets shall never appear in client bundles or logs. Credentials and private account data shall not be sent to AI providers. Case text shall be shared only according to the user disclosure and consent process. Provider failures must leave booking and account data intact.')
h('Deployment boundary')
p('Local development and a controlled pilot are the first environments. Public deployment requires access controls, a configured retention policy, tested language capabilities, and the release acceptance checks. These requirements do not authorize real payment collection or external notifications.')

# 5
page('4 Accounts and advocate onboarding')
for code,t in [
 ('FR AUTH 01','The system SHALL register user and advocate accounts with a unique normalized email, name, password and preferred language. Administrator creation SHALL use a protected bootstrap process.'),
 ('FR AUTH 02','The system SHALL validate fields and store passwords using an established password-hashing library with a maintained configuration. It SHALL never store or return plaintext passwords.'),
 ('FR AUTH 03','The system SHALL authenticate sessions, support logout and expiry, and invalidate access when an account is disabled. Session handling SHALL follow the security controls in section 10.'),
 ('FR AUTH 04','The backend SHALL derive account identity and role from the authenticated session. Client-provided user IDs or roles SHALL not grant access or ownership.'),
 ('FR AUTH 05','Every private operation SHALL check role, ownership and assignment. Direct requests bypassing interface restrictions SHALL receive the same access checks.'),
 ('FR ADV 01','Advocates SHALL submit expertise, languages, location, experience, consultation mode, professional enrollment details and supporting credentials.'),
 ('FR ADV 02','New applications SHALL start pending. Only an authorized administrator SHALL approve, reject or suspend an application, with a reason and audit record.'),
 ('FR ADV 03','Public discovery SHALL include only approved, active profiles. Private credential files SHALL never appear in public profile responses.'),
 ('FR ADV 04','Advocates SHALL edit their own profiles. Changes to credentials or enrollment identity SHALL trigger review and pause public discovery until approval.'),
 ('FR ADV 05','The administrator SHALL see a minimal review queue and decision history. A verification label SHALL describe the checks completed, without promising legal outcomes.')]:req(code,t)
h('Application state')
p('Allowed application states are pending, approved, rejected and suspended. A rejected applicant may resubmit to pending. Suspension prevents new consultation requests. Existing future bookings require explicit handling; the system must not silently cancel or conceal them.')
h('Acceptance evidence')
p('Demonstrate two independent accounts, a pending advocate hidden from search, an administrator approval that publishes the profile, and rejected attempts to access another account’s private records. Password reset and external identity verification are P1 unless required by the selected deployment policy.')

# 6
page('5 Multilingual intake and triage')
for code,t in [
 ('FR LNG 01','The system SHALL maintain a language registry with stable locale codes and independent capability flags for interface, text response, recognition and playback.'),
 ('FR LNG 02','The user SHALL explicitly select a language. The system SHALL preserve that choice through intake, follow-ups and results, rather than mapping all non-English input to Hindi.'),
 ('FR LNG 03','Unsupported recognition or playback SHALL show a clear explanation and alternative. Speech failure SHALL not prevent typed submission or erase reviewed text.'),
 ('FR INP 01','Users SHALL type an issue or capture voice through a supported adapter. Microphone access SHALL require browser permission and a visible recording state.'),
 ('FR INP 02','Users SHALL review and edit the transcript before analysis. The system SHALL not submit interim recognition results automatically.'),
 ('FR TRI 01','Triage SHALL accept reviewed text, language and authorized case context. The server SHALL limit input size and validate all fields before calling a provider.'),
 ('FR TRI 02','The result SHALL include domain, urgency, a short summary, suggested actions, response language, source mode and limitations. Domains SHALL cover criminal, civil, family and property, with an unknown state for unresolved input.'),
 ('FR TRI 03','The backend SHALL validate result types and permitted values. Invalid provider output SHALL trigger an explicit failure or labelled limited fallback; it SHALL not silently default an unknown issue to civil.'),
 ('FR TRI 04','Follow-ups SHALL use only the authorized case’s recent context, with a configurable bound. The system SHALL prevent mixing unrelated cases or users.'),
 ('FR TRI 05','Every AI result SHALL state that it provides preliminary information. Statutory citations and deadlines SHALL be omitted unless checked against reviewed sources. Urgency SHALL not be presented as a verified professional assessment.')]:req(code,t)
h('Proposed input contract limits')
p('Initial configuration: 10,000 characters per message and the most recent 10 conversation messages per triage call. The interface must communicate exceeded limits. Provider timeout defaults to 30 seconds; the user receives a retry or a clearly limited result when it expires. These are design limits, not measured performance results.')
h('Language evaluation')
p('Evaluate names, dates, amounts, code-switching and common accents in each pilot language. Record recognition and playback support separately. A response in an unexpected language shall be flagged and retried or shown with an explicit limitation.')

# 7
page('6 Advocate discovery and private summaries')
for code,t in [
 ('FR MAT 01','Matching SHALL filter approved, active advocates by the selected practice area and requested communication language. A user SHALL be able to correct the suggested area.'),
 ('FR MAT 02','Location filters SHALL use the advocate’s published city or district. The user SHALL be able to broaden location explicitly; the system SHALL not weaken a required language filter silently.'),
 ('FR MAT 03','Results SHALL indicate relevant expertise, language suitability, location and published slot availability. Ranking SHALL use documented relevance rules, not fabricated ratings or outcome claims.'),
 ('FR MAT 04','Discovery SHALL support pagination, filter reset, and empty results. Only future, unreserved, valid slots SHALL count as requestable availability.'),
 ('FR PRO 01','Public profiles SHALL show name, expertise, languages, location, experience, consultation mode, approval explanation and available slots. Contact sharing SHALL follow the confirmed booking flow.'),
 ('FR CAS 01','The system SHALL draft an editable intake summary containing confirmed facts, location, key dates and questions. Missing information SHALL remain missing or explicitly unresolved.'),
 ('FR CAS 02','Users SHALL explicitly approve sharing a specific summary version with a selected advocate. The system SHALL store the recipient, version and consent timestamp.'),
 ('FR CAS 03','Advocates SHALL access only summaries shared with them for assigned requests. Raw transcripts and private notes SHALL not become public or be shared automatically.'),
 ('FR CAS 04','Users SHALL view their own saved cases and delete or withdraw unshared draft information. Deletion of shared or retained records SHALL follow the configured retention policy and be explained.')]:req(code,t)
h('Matching rules')
p('First apply approval, active account, expertise and language constraints. Apply location when the user requires it. Order remaining results by the documented expertise match and earliest requestable slot, using a stable tie-breaker. Report why each profile matches. Do not promise that any profile is the best advocate or will respond immediately.')
h('Privacy boundary')
p('Cases and summaries are private by default. Sharing a summary does not authorize publication, provider training use, or unrestricted administrator access. A later change to the case shall not overwrite the version already approved for a booking; the user must approve a new shared version.')

# 8
page('7 Slot reservation and booking lifecycle')
for code,t in [
 ('FR SLT 01','Approved advocates SHALL publish future slots with start/end times and a consultation mode. Times SHALL be stored in UTC and displayed with an explicit user-facing timezone.'),
 ('FR SLT 02','The system SHALL reject invalid durations, past slots and overlapping active slots for the same advocate. Reservation and slot publication SHALL remain safe under concurrent requests.'),
 ('FR BKG 01','An authenticated user SHALL request a slot using its ID and approved summary version. The server SHALL derive the user, advocate, times and mode from trusted records.'),
 ('FR BKG 02','Creating a request SHALL atomically reserve the slot. At most one active requested or accepted booking SHALL exist for a slot.'),
 ('FR BKG 03','An assigned approved advocate SHALL accept or decline a pending request. The system SHALL reject unauthorized, expired or repeated conflicting transitions.'),
 ('FR BKG 04','Requested bookings SHALL expire at the earlier of slot start or 24 hours after creation. The backend SHALL enforce expiry and release slots independently of an open browser.'),
 ('FR BKG 05','Users SHALL cancel their own future requests or accepted appointments. Advocates SHALL provide a reason when cancelling assigned future appointments.'),
 ('FR BKG 06','Booking lists SHALL persist across logout and restart, with status history visible to authorized participants. Refreshing or retrying SHALL not create duplicate requests.'),
 ('FR BKG 07','The system SHALL notify affected participants in-app after a status change. A booking SHALL remain correct if notification delivery fails.')]:req(code,t)
table(['State','Allowed next state','Slot effect'],[
 ['requested','accepted, declined, cancelled, expired','Reserved until terminal decision or expiry'],
 ['accepted','cancelled, completed','Reserved; completion permitted only after slot end'],
 ['declined / cancelled / expired','None; a new request gets a new ID','Release only if future and advocate remains eligible'],
 ['completed','None','Historical record; no new request for the past slot']], [1.4,2.7,2.73],9)
p('Completion is recorded by the assigned advocate after the scheduled end, with an audit event. Rescheduling is deferred: cancel and submit a new request. This release collects no payment, and acceptance must not depend on a demo payment bypass.')

# 9
page('8 External interfaces and API contracts')
p('The following routes form a proposed contract for the new build. Developers may revise names jointly before implementation, but must preserve semantics, permissions and acceptance coverage. All protected routes require a valid authenticated session.')
table(['Method and path','Actor','Required behavior'],[
 ['POST /api/auth/register','Visitor','Create user or advocate; reject administrator role'],
 ['POST /api/auth/login\nPOST /api/auth/logout','Account','Start / revoke session'],
 ['GET /api/auth/me','Account','Return safe identity and permissions'],
 ['PUT /api/advocates/me','Advocate','Save own application/profile'],
 ['GET /api/admin/advocates\nPATCH /api/admin/advocates/:id/status','Admin','Review and record approval decision'],
 ['GET /api/languages','Public','Return language capability registry'],
 ['POST /api/ai/triage','User','Reviewed text, locale and own case context'],
 ['GET /api/advocates\nGET /api/advocates/:id','Public','Approved profiles and permitted public fields'],
 ['POST /api/cases\nGET /api/cases\nGET /api/cases/:id\nPATCH /api/cases/:id\nDELETE /api/cases/:id','User','Create, view, edit or delete own permitted case information'],
 ['POST /api/cases/:id/consent','User','Approve version and recipient'],
 ['GET /api/advocates/:id/slots\nPOST /api/slots\nDELETE /api/slots/:id','Public / advocate','Read availability; publish/remove own free slots'],
 ['POST /api/bookings\nGET /api/bookings','User / advocate','Reserve eligible slot; list own/assigned records'],
 ['PATCH /api/bookings/:id/status','Participant','Apply an authorized lifecycle transition'],
 ['GET /api/notifications\nPATCH /api/notifications/:id/read','Account','List own notifications and mark read'],
 ['POST /api/advocates/me/credentials\nGET /api/admin/credentials/:id','Advocate / admin','Upload own credentials; authorized private review']], [3.2,1.0,2.63],8.7)
h('Contract conventions')
p('Use JSON for ordinary requests. Successful responses use data plus optional pagination metadata. Errors use code, message, fieldErrors and requestId. Return 400 for invalid input, 401 for missing/expired session, 403 for forbidden actions, 404 for inaccessible resources, 409 for conflicts, 429 for rate limits, and 503 for provider unavailability without an enabled fallback.')
p('Collections default to 20 records, maximum 50. Use ISO 8601 UTC timestamps. Booking creation accepts slotId, caseId, summaryVersion and an idempotency key; identity derives from the session. Triage uses locale codes, domain values criminal/civil/family/property/unknown, urgency low/medium/high/emergency/unknown, and sourceMode provider/fallback. Credential uploads use protected multipart requests.')

# 10
page('9 Data requirements and persistence')
table(['Entity','Required fields','Integrity and access'],[
 ['Account','ID, name, normalized email, password hash, role, preferred locale, disabled flag, timestamps','Unique email; role constrained; hashes never serialized'],
 ['Advocate application','Account ID, enrollment details, credential references, review state, decision reason/history','One application per advocate; private credentials; audited changes'],
 ['Public profile','Advocate ID, expertise/tags, languages, city/district, experience, consultation modes','Published only for approved active accounts'],
 ['Case and summary','Owner ID, reviewed text, locale, bounded conversation, result, summary versions, consent events','Private by default; version-specific recipient access'],
 ['Consultation slot','Advocate ID, UTC start/end, mode, reservation reference, active flag','No overlapping active slots; exclusive reservation'],
 ['Booking','User ID, advocate ID, slot ID, case ID, shared version, state, expiry, transition history','One active booking per slot; idempotent creation'],
 ['Notification and audit','Recipient/actor ID, event type, record reference, timestamp, read state or decision','Restrict recipient access; exclude private content from logs']], [1.25,3.0,2.58],9.5)
h('Consistency requirements')
req('DR 01','The database SHALL persist accounts, applications, cases, slots and bookings across application restarts. Mock objects or browser storage alone SHALL not satisfy release acceptance.')
req('DR 02','Unique constraints and atomic reservation logic SHALL enforce exclusive booking. Slot-overlap creation SHALL use a concurrency-safe strategy, such as serialized per-advocate schedule updates, rather than an unprotected read-then-write check.')
req('DR 03','Booking transitions and reservation release SHALL be committed consistently. Failed requests SHALL not leave orphan reservations. Retried operations SHALL return the existing result or a defined conflict.')
req('DR 04','Indexes SHALL cover email uniqueness, advocate approval/expertise/languages/location, slot lookup by advocate/time, and booking lookup by user/advocate/status. Migration and fixture scripts SHALL be repeatable.')
h('Retention and uploads')
p('The deployment must define retention periods before collecting real case content or credentials. Draft deletion must remove the associated data according to that policy, and users must understand retained booking/audit records. Permit credential PDF/JPEG/PNG uploads up to 5 MB per file, validate type/content, randomize storage names, and prohibit public URLs.')

# 11
page('10 Nonfunctional and security requirements')
for code,t in [
 ('NFR SEC 01','The system SHALL enforce HTTPS for a public pilot, secure password hashing, restricted CORS and server-side authorization. Cookie sessions SHALL use HttpOnly and appropriate Secure/SameSite settings with CSRF protection for state-changing requests.'),
 ('NFR SEC 02','Authentication and expensive AI endpoints SHALL have configurable rate limits. Validate IDs, enums, lengths, uploads and database queries. Return safe errors without stack traces or secrets.'),
 ('NFR SEC 03','Logs SHALL record request IDs, decision events and failures without plaintext passwords, credentials, full case text or session tokens. Private records SHALL be excluded from public serializers.'),
 ('NFR SEC 04','Session expiry, logout, disabled accounts and privilege changes SHALL take effect on backend access checks. Sensitive credential retrieval SHALL be limited to authorized reviewers and audited.'),
 ('NFR PRI 01','Before provider submission or advocate sharing, the interface SHALL explain the intended recipient and purpose. Collect only necessary information and apply the configured retention/deletion policy.'),
 ('NFR REL 01','Database failure SHALL produce an explicit unavailable state. The system SHALL not display an unpersisted request as a successful booking. Notification or provider failure SHALL not corrupt appointments.'),
 ('NFR PERF 01','Proposed pilot target: at 20 concurrent sessions, ordinary API operations SHOULD have p95 response time under 1 second, excluding external providers and uploads. Measure in a documented environment; this is not an existing benchmark.'),
 ('NFR PERF 02','The interface SHALL display progress during triage, prevent duplicate submission and allow an intelligible timeout/retry path. Provider timeout SHALL be configurable, initially 30 seconds.'),
 ('NFR UX 01','Core flows SHALL work at 360 px mobile width and desktop widths, with keyboard access, visible focus, labelled inputs, readable contrast, and non-color-only status indicators.'),
 ('NFR LNG 01','Text SHALL preserve Unicode and relevant scripts. Speech controls SHALL expose availability and failure states. Testing SHALL include the selected pilot languages and unsupported-mode fallbacks.'),
 ('NFR OPS 01','The team SHALL provide environment examples, a health/readiness check, deterministic development setup, backups and a documented restore test for the pilot database.'),
 ('NFR MAINT 01','Provider adapters, language registry, permissions and booking transitions SHALL have separated responsibilities and meaningful tests. The API contract SHALL remain versioned alongside implementation.')]:req(code,t)
p('Legal/content correctness requires professional review in addition to software testing. Structured JSON, prompts, or local rules do not establish legal accuracy. The system must never promise case outcomes, universal language quality, or guaranteed urgent response.')

# 12
page('11 Acceptance tests and traceability')
p('A working release must pass the following checks against real frontend/backend integration and persistent records. Development mocks cannot satisfy these criteria. The IDs link each test to its requirements.')
table(['Test','Scenario and expected result','Requirements'],[
 ['AT 01','Register two accounts, log in/out and expire a session. Unauthorized calls fail; passwords are absent from responses.','AUTH 01–05; SEC 01–04'],
 ['AT 02','Pending advocate stays hidden. Administrator approval publishes profile; suspended advocate cannot receive new requests.','ADV 01–05; MAT 01'],
 ['AT 03','Review voice transcript, edit a fact, submit typed fallback when speech is unsupported. Selected locale persists.','LNG 01–03; INP 01–02'],
 ['AT 04','Provider failure/invalid JSON produces a labelled limitation or error. Follow-up context never leaks between users/cases.','TRI 01–05; REL 01'],
 ['AT 05','Match approved advocates by expertise and language. Empty results are explicit and location broadening requires user action.','MAT 01–04; PRO 01'],
 ['AT 06','User approves a summary version. Only assigned recipient can view it; later edits do not overwrite shared version.','CAS 01–04; PRI 01'],
 ['AT 07','Publish future slot; reject past or overlapping slots. Two simultaneous requests yield exactly one reservation.','SLT 01–02; BKG 01–02; DR 02'],
 ['AT 08','Assigned advocate accepts/declines. Other accounts cannot change it. User sees persisted status after restart.','BKG 03, 06; AUTH 05; DR 01'],
 ['AT 09','Expiry and cancellation release eligible slots. Repeated requests/transitions do not duplicate or corrupt records.','BKG 04–07; DR 03'],
 ['AT 10','Run core journey on mobile/keyboard and pilot languages; verify timeout, rate-limit and private-file handling.','NFR UX 01; LNG 01; SEC 02–04']], [0.63,4.8,1.4],9)
h('Final demonstration')
p('Use separate user, advocate and administrator accounts. Approve an advocate, publish a slot, submit an issue, select a suitable profile, approve a summary, request the slot, and accept it from the advocate account. Confirm accepted status from the user account after refresh/restart. Attempt the same slot from a second user and prove the conflict is safe.')
h('Evaluation beyond software tests')
p('Prepare an advocate-reviewed case set and language-specific comprehension trials. Record sample size, cases, scripts, failures and limitations. Thresholds for domain/urgency quality must be agreed with reviewers before claiming accuracy. The release must report actual tested coverage and must not invent performance or social-impact results.')

# 13
page('12 Delivery plan and open decisions')
h('Implementation order')
table(['Stage','Backend owner','Frontend owner','Exit condition'],[
 ['1 Contract and setup','Data/state model; API contract; fixtures','Screen flow and matching mock payloads','Both owners agree on contracts'],
 ['2 Accounts and profiles','Authentication; approval; public serialization','Registration; role dashboards; approval screen','Approved advocate appears to user'],
 ['3 Intake and discovery','Triage adapter; capabilities; filters','Transcript review; language selection; results','Reviewed issue reaches suitable profiles'],
 ['4 Slots and requests','Atomic reservation; lifecycle; expiry','Slot selection; request list; acceptance UI','User request becomes accepted'],
 ['5 Integration and hardening','Security/concurrency tests; persistence','Failure/mobile/accessibility checks','All release acceptance checks pass']], [0.9,2.0,2.0,1.93],9)
h('Open decisions before a real pilot')
p('Confirm pilot languages and provider capabilities; the AI provider and data handling terms; administrator review criteria; account recovery policy; case/credential retention periods; consultation contact-sharing method; and hosting configuration. These are launch decisions, not evidence of completed capabilities. The proposed 24-hour request expiry and performance targets may be revised through the change process.')
h('Definition of done')
p('Deliver runnable source, installation instructions, environment examples, repeatable migrations/fixtures, current API contracts, and test results. The complete journey must operate on persistent data. All sample profiles must be labelled. No P0 defect may remain in permissions, reservation, acceptance or status persistence. Document external dependencies and incomplete language capabilities explicitly.')
h('Change control')
p('Record requirement changes with the affected ID, reason, owner, acceptance impact and version date. Both developers must review incompatible contract changes. Deferred payments and other extensions require a separate approved scope and tests; their existence in a mockup does not make them part of this release.')
h('Basis and terminology')
p('This document derives its scope from the team’s request for a new build, Indian multilingual access, separate user/advocate roles, and a time-constrained working product. Technology choices and numeric limits are proposed design decisions. The specification does not incorporate old code as an implementation dependency or claim formal standards certification.')
p('Glossary: MVP means minimum viable product; triage means preliminary categorization; locale means a stable language identifier; p95 means the 95th percentile; idempotency means repeated requests do not create duplicate effects; approval means completion of the defined administrator review.')

# Remove template title rules and theme borders so the black title stands alone.
for el in list(d.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for el in list(d.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)
print(OUT)

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# COLOR PALETTE DEFINITION (Modern Corporate Legal Tech)
# ==============================================================================
HEX_NAVY = "0B2545"        # Primary Deep Navy
HEX_BLUE = "1D4ED8"        # Royal Accent Blue
HEX_DARK = "0F172A"        # Slate Dark / Heading
HEX_BODY = "334155"        # Slate Body Text
HEX_MUTED = "64748B"       # Slate Muted
HEX_BORDER = "CBD5E1"      # Slate Light Border
HEX_BG_ROW = "F8FAFC"      # Light Table Row
HEX_AMBER = "B45309"       # Warning / Statutory Gold
HEX_AMBER_BG = "FEF3C7"    # Amber Light
HEX_GREEN = "166534"       # Success / Practical Tip
HEX_GREEN_BG = "F0FDF4"    # Green Light
HEX_PURPLE = "6B21A8"      # Viva Defense Accent
HEX_PURPLE_BG = "FAF5FF"   # Purple Light
HEX_BLUE_BG = "EFF6FF"     # Info / Architecture Light
HEX_CODE_BG = "F1F5F9"     # Monospace Block Light

COLOR_NAVY = RGBColor(11, 37, 69)
COLOR_BLUE = RGBColor(29, 78, 216)
COLOR_DARK = RGBColor(15, 23, 42)
COLOR_BODY = RGBColor(51, 65, 85)
COLOR_MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_left_border(cell, hex_color=HEX_BLUE, sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLUE
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK
    return p

def add_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_BODY
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_BODY
    return p

def add_callout(doc, title, text, style_type="note"):
    cfg = {
        "note": {"border": HEX_BLUE, "bg": HEX_BLUE_BG, "title_color": RGBColor(29, 78, 216)},
        "statute": {"border": HEX_AMBER, "bg": HEX_AMBER_BG, "title_color": RGBColor(180, 83, 9)},
        "tip": {"border": HEX_GREEN, "bg": HEX_GREEN_BG, "title_color": RGBColor(22, 101, 52)},
        "viva": {"border": HEX_PURPLE, "bg": HEX_PURPLE_BG, "title_color": RGBColor(107, 33, 168)}
    }.get(style_type, {"border": HEX_BLUE, "bg": HEX_BLUE_BG, "title_color": RGBColor(29, 78, 216)})

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]
    set_cell_background(c, cfg["bg"])
    set_cell_margins(c, top=90, bottom=90, left=140, right=140)
    set_cell_left_border(c, hex_color=cfg["border"], sz="24")

    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_t = p.add_run(title.upper() + "\n")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    r_t.font.color.rgb = cfg["title_color"]

    r_b = p.add_run(text)
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(9)
    r_b.font.color.rgb = COLOR_BODY

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]
    set_cell_background(c, HEX_CODE_BG)
    set_cell_margins(c, top=70, bottom=70, left=100, right=100)
    
    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def create_styled_table(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, HEX_BORDER)

    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        c = hdr_cells[i]
        set_cell_background(c, HEX_NAVY)
        set_cell_margins(c, top=90, bottom=90, left=110, right=110)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = HEX_BG_ROW if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            c = row_cells[c_idx]
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=70, bottom=70, left=110, right=110)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(cell_value))
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_BODY

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

print("Ready to construct document content.")
